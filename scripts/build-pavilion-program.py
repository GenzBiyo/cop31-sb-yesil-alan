# -*- coding: utf-8 -*-
"""Turn the ministry programme workbook into src/lib/pavilion-program.ts."""
import json
import re
from datetime import datetime, date
from pathlib import Path

import openpyxl

XLSX = Path(r"c:\Users\umut.agyuz\Downloads\COP31_Saglik_Bakanligi_Yesil_Alan_Programi (4).xlsx")
OUT = Path(__file__).resolve().parents[1] / "src" / "lib" / "pavilion-program.ts"

THEMES = {
    "2026-11-09": {"tr": "Gıda, tarım ve sağlık", "en": "Food, Agriculture and Health"},
    "2026-11-10": {"tr": "Enerji ve ulaştırma", "en": "Energy and Transport"},
    "2026-11-11": {"tr": "Sıfır Atık", "en": "Zero Waste"},
    "2026-11-12": {"tr": "Dayanıklı şehirler ve yapılı çevre", "en": "Resilient Cities and Built Environment"},
    "2026-11-13": {"tr": "Finans ve ticaret", "en": "Finance and Trade"},
    "2026-11-14": {"tr": "Çocuklar, gençlik, eğitim ve beceriler", "en": "Children, Youth, Education and Skills"},
    "2026-11-15": {"tr": "Antalya'da bir nefes", "en": "A Breath in Antalya"},
    "2026-11-16": {"tr": "Bilim, sanayi ve teknoloji", "en": "Science, Industry and Technology"},
    "2026-11-17": {"tr": "Okyanus, denizler, doğa ve arazi kullanımı — Rio Sinerjisi", "en": "Ocean, Nature and Land Use — Rio Synergies"},
    "2026-11-18": {"tr": "İnsani ve sosyal kalkınma", "en": "Humanitarian and Social Development"},
    "2026-11-19": {"tr": "İmece: uygulamanın güçlendirilmesi", "en": "İmece: Strengthening Implementation"},
    "2026-11-20": {"tr": "Nihai müzakereler", "en": "Final Negotiations"},
}


def norm(value: str) -> str:
    text = (value or "").lower().replace("–", " ").replace("—", " ").replace("-", " ")
    text = re.sub(r"\[.*?\]", " ", text)
    text = re.sub(r"[^a-z0-9çğıöşü ]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def tokens(value: str):
    stop = {"ve", "ile", "icin", "için", "bir", "the", "of", "and", "in", "için"}
    return [part for part in norm(value).split() if len(part) > 2 and part not in stop]


def cell_parts(value):
    parts = [part.strip() for part in str(value or "").split("\n")]
    parts = [part for part in parts if part]
    title = re.sub(r"\s*\[.*?\]\s*", "", parts[0]).strip(" -–|") if parts else ""
    status = ""
    host = ""
    for part in parts[1:]:
        if part.startswith("["):
            status = part.strip("[]")
        elif not host:
            host = part
    return title, host, status


def clean_speaker(name: str) -> str:
    text = name.strip()
    if text.upper() in {"SANOFI", "SANOFİ"}:
        return "Sanofi Türkiye"
    return text


def js(value) -> str:
    return json.dumps(value, ensure_ascii=False)


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    programs = []
    for row in list(wb["Program"].iter_rows(values_only=True))[4:]:
        if not row or not row[0] or not row[5]:
            continue
        raw_date = row[1]
        if isinstance(raw_date, datetime):
            day = raw_date.date().isoformat()
        else:
            day = str(raw_date or "")[:10]
        speakers = []
        for cell in row[10:14]:
            name = clean_speaker(str(cell or ""))
            if name and name not in speakers:
                speakers.append(name)
        programs.append(
            {
                "code": str(row[0]).strip(),
                "date": day,
                "slot": str(row[3] or ""),
                "theme": str(row[4] or "").strip(),
                "title": str(row[5] or "").strip(),
                "titleEn": str(row[6] or "").strip(),
                "kind": str(row[7] or "").strip(),
                "host": str(row[8] or "").strip(),
                "moderator": str(row[9] or "").strip(),
                "speakers": speakers,
                "topics": str(row[15] or "").strip() if row[15] else "",
                "status": str(row[17] or "").strip() if len(row) > 17 and row[17] else "",
                "note": str(row[18] or "").strip() if len(row) > 18 and row[18] else "",
            }
        )

    cal = wb["Bir Bakışta Takvim"]
    slots = [("10:00", "11:00", 14), ("14:00", "15:00", 15), ("16:00", "17:00", 16)]
    days = [date(2026, 11, 9 + index).isoformat() for index in range(12)]
    calendar = []
    for start, end, row_index in slots:
        for index, day in enumerate(days):
            value = cal.cell(row_index, index + 2).value
            if not value:
                continue
            title, host, status = cell_parts(value)
            if not title:
                continue
            calendar.append(
                {"date": day, "start": start, "end": end, "title": title, "host": host, "status": status}
            )

    used = set()
    sessions = []
    for item in calendar:
        if "kapalı" in item["title"].lower() or "kapali" in norm(item["title"]):
            continue
        best = None
        best_score = 0
        left = tokens(item["title"])
        for index, program in enumerate(programs):
            if index in used:
                continue
            right = tokens(program["title"])
            if not left or not right:
                continue
            overlap = len(set(left) & set(right))
            score = overlap / max(1, min(len(set(left)), len(set(right))))
            prefix = norm(item["title"])[:28]
            if prefix and prefix in norm(program["title"]):
                score += 0.45
            if score > best_score:
                best_score = score
                best = index
        program = programs[best] if best is not None and best_score >= 0.55 else None
        if program is None and "karbon ayak" in norm(item["title"]):
            for index, candidate in enumerate(programs):
                if candidate["code"] == "10-S1" and index not in used:
                    program = candidate
                    best = index
                    break
        if program:
            used.add(best)
            if program["code"] == "10-S1":
                program = {**program, "title": item["title"], "titleEn": program["titleEn"]}
        status = (program or {}).get("status") or item["status"] or "Planlandı"
        sessions.append(
            {
                "code": (program or {}).get("code") or f"{item['date']}-{item['start']}",
                "date": item["date"],
                "start": item["start"],
                "end": item["end"],
                "title": (program or {}).get("title") or item["title"],
                "titleEn": (program or {}).get("titleEn") or "",
                "theme": THEMES[item["date"]]["tr"],
                "kind": (program or {}).get("kind") or "Panel",
                "host": (program or {}).get("host") or item["host"],
                "moderator": (program or {}).get("moderator") or "",
                "speakers": (program or {}).get("speakers") or [],
                "topics": (program or {}).get("topics") or "",
                "status": status,
                "note": (program or {}).get("note") or "",
            }
        )

    for index, program in enumerate(programs):
        if program["code"] == "13-MA" and index not in used:
            used.add(index)
            sessions.append(
                {
                    "code": program["code"],
                    "date": program["date"],
                    "start": "12:00",
                    "end": "13:00",
                    "title": program["title"],
                    "titleEn": program["titleEn"],
                    "theme": THEMES.get(program["date"], {}).get("tr", ""),
                    "kind": program["kind"],
                    "host": program["host"],
                    "moderator": program["moderator"],
                    "speakers": program["speakers"],
                    "topics": program["topics"],
                    "status": program["status"] or "Teyit bekliyor",
                    "note": program["note"],
                }
            )

    sessions.sort(key=lambda item: (item["date"], item["start"], item["code"]))

    stands = []
    sheet = wb["Stantlar"]
    blocks = [("10:00", "11:00"), ("14:00", "15:00"), ("16:00", "17:00")]
    for row in list(sheet.iter_rows(values_only=True))[5:]:
        raw = row[0]
        if isinstance(raw, datetime):
            day = raw.date().isoformat()
        else:
            continue
        for offset, label in ((2, "Stant 1"), (3, "Stant 2"), (4, "Stant 3")):
            name = str(row[offset] or "").strip()
            if name:
                stands.append({"date": day, "stand": label, "name": name, "start": "10:00", "end": "18:00"})
        firm_cols = [(5, "Stant 4"), (8, "Stant 5"), (11, "Stant 6")]
        for start_col, label in firm_cols:
            for block, (start, end) in enumerate(blocks):
                name = str(row[start_col + block] or "").strip()
                if not name or name in {"?", "#REF!"}:
                    continue
                stands.append({"date": day, "stand": label, "name": name, "start": start, "end": end})

    lines = [
        "/** Generated from COP31 Yeşil Alan program workbook. Edit the workbook, then regenerate. */",
        "export type PavilionSession = {",
        "  code: string;",
        "  date: string;",
        "  start: string;",
        "  end: string;",
        "  title: string;",
        "  titleEn: string;",
        "  theme: string;",
        "  kind: string;",
        "  host: string;",
        "  moderator: string;",
        "  speakers: string[];",
        "  topics: string;",
        "  status: string;",
        "  note: string;",
        "};",
        "",
        "export type PavilionStand = { date: string; stand: string; name: string; start: string; end: string };",
        "",
        "export const PAVILION_THEMES: Record<string, { tr: string; en: string }> = " + js(THEMES) + ";",
        "",
        "export const PAVILION_CLOSED = [\"2026-11-11\"];",
        "",
        "export const PAVILION_SESSIONS: PavilionSession[] = " + js(sessions) + ";",
        "",
        "export const PAVILION_STANDS: PavilionStand[] = " + js(stands) + ";",
        "",
    ]
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"sessions {len(sessions)} stands {len(stands)} -> {OUT}")
    for session in sessions:
        print(f"{session['date']} {session['start']} {session['code']} {session['status']} {session['title'][:68]}")
    print("--- unused ---")
    for index, program in enumerate(programs):
        if index not in used:
            print(program["code"], program["title"][:70])


if __name__ == "__main__":
    main()
